from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.schemas.cover_letter import CoverLetter


class CoverLetterDocumentGenerator:

    def generate(
        self,
        cover_letter: CoverLetter,
        application_id: int,
    ) -> str:

        output_dir = Path(
            f"generated/application_{application_id}"
        )

        output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_path = output_dir / "cover_letter.docx"

        document = Document()

        # -------------------------------------------------
        # PAGE SETTINGS
        # -------------------------------------------------

        section = document.sections[0]

        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # -------------------------------------------------
        # DEFAULT FONT
        # -------------------------------------------------

        normal_style = document.styles["Normal"]

        normal_style.font.name = "Arial"
        normal_style.font.size = Pt(10.5)

        paragraph_format = normal_style.paragraph_format
        paragraph_format.space_after = Pt(8)
        paragraph_format.line_spacing = 1.08

        # -------------------------------------------------
        # CANDIDATE NAME
        # -------------------------------------------------

        name_paragraph = document.add_paragraph()

        name_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        name_run = name_paragraph.add_run(
            cover_letter.candidate_name.upper()
        )

        name_run.bold = True
        name_run.font.name = "Arial"
        name_run.font.size = Pt(17)

        # -------------------------------------------------
        # SUBJECT
        # -------------------------------------------------

        if cover_letter.subject:
            subject_paragraph = document.add_paragraph()

            subject_run = subject_paragraph.add_run(
                f"Konu: {cover_letter.subject}"
                if cover_letter.language == "tr"
                else f"Subject: {cover_letter.subject}"
            )

            subject_run.bold = True
            subject_run.font.name = "Arial"
            subject_run.font.size = Pt(10.5)

        # -------------------------------------------------
        # GREETING
        # -------------------------------------------------

        greeting = document.add_paragraph(
            cover_letter.greeting
        )

        greeting.paragraph_format.space_after = Pt(10)

        # -------------------------------------------------
        # OPENING
        # -------------------------------------------------

        opening = document.add_paragraph(
            cover_letter.opening
        )

        opening.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # -------------------------------------------------
        # BODY
        # -------------------------------------------------

        for body_text in cover_letter.body_paragraphs:

            paragraph = document.add_paragraph(
                body_text
            )

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
            )

        # -------------------------------------------------
        # CLOSING
        # -------------------------------------------------

        closing = document.add_paragraph(
            cover_letter.closing
        )

        closing.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        closing.paragraph_format.space_before = Pt(4)

        # -------------------------------------------------
        # SIGN OFF
        # -------------------------------------------------

        sign_off = document.add_paragraph()

        sign_off.paragraph_format.space_before = Pt(12)
        sign_off.paragraph_format.space_after = Pt(2)

        sign_off_run = sign_off.add_run(
            cover_letter.sign_off
        )

        sign_off_run.font.name = "Arial"

        candidate = document.add_paragraph()

        candidate.paragraph_format.space_after = Pt(0)

        candidate_run = candidate.add_run(
            cover_letter.candidate_name
        )

        candidate_run.bold = True
        candidate_run.font.name = "Arial"

        # -------------------------------------------------
        # SAVE
        # -------------------------------------------------

        document.save(output_path)

        return str(output_path)


cover_letter_document_generator = (
    CoverLetterDocumentGenerator()
)