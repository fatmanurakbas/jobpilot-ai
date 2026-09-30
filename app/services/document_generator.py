from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.schemas.candidate_profile import CandidateProfile
from app.schemas.cv_tailoring import CVTailoringPlan


class DocumentGenerator:

    BASE_OUTPUT_DIR = Path("generated")

    # ---------------------------------------------------------
    # SECTION TITLES
    # ---------------------------------------------------------

    SECTION_TITLES = {
        "tr": {
            "summary": "Profesyonel Özet",
            "skills": "Teknik Yetenekler",
            "experience": "Deneyim",
            "projects": "Projeler",
            "education": "Eğitim",
            "certifications": "Sertifikalar",
            "languages": "Diller",
        },
        "en": {
            "summary": "Professional Summary",
            "skills": "Technical Skills",
            "experience": "Experience",
            "projects": "Projects",
            "education": "Education",
            "certifications": "Certifications",
            "languages": "Languages",
        },
    }

    # ---------------------------------------------------------
    # MAIN GENERATOR
    # ---------------------------------------------------------

    def generate_tailored_cv(
        self,
        candidate: CandidateProfile,
        tailoring_plan: CVTailoringPlan,
        application_id: int,
    ) -> str:

        output_dir = (
            self.BASE_OUTPUT_DIR
            / f"application_{application_id}"
        )

        output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_path = (
            output_dir
            / "tailored_cv.docx"
        )

        document = Document()

        self._configure_document(document)

        self._add_header(
            document,
            candidate,
        )

        self._add_summary(
            document,
            candidate,
            tailoring_plan,
        )
        
        self._add_education(
            document,
            candidate,
        )

        self._add_skills(
            document,
            candidate,
            tailoring_plan,
        )

        self._add_experience(
            document,
            candidate,
            tailoring_plan,
        )

        self._add_projects(
            document,
            candidate,
            tailoring_plan,
        )

        self._add_certifications(
            document,
            candidate,
        )

        self._add_languages(
            document,
            candidate,
        )

        document.save(output_path)

        return str(output_path)

    # ---------------------------------------------------------
    # DOCUMENT SETTINGS
    # ---------------------------------------------------------

    def _configure_document(
        self,
        document: Document,
    ) -> None:

        section = document.sections[0]

        # Compact but readable ATS-friendly margins
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

        normal_style = document.styles["Normal"]

        normal_style.font.name = "Arial"
        normal_style.font.size = Pt(9.5)

        normal_style.paragraph_format.space_before = Pt(0)
        normal_style.paragraph_format.space_after = Pt(2)
        normal_style.paragraph_format.line_spacing = 1.05

    # ---------------------------------------------------------
    # LANGUAGE
    # ---------------------------------------------------------

    def _get_language(
        self,
        candidate: CandidateProfile,
    ) -> str:

        language = (
            candidate.cv_language or "en"
        ).lower()

        if language not in self.SECTION_TITLES:
            return "en"

        return language

    def _title(
        self,
        candidate: CandidateProfile,
        key: str,
    ) -> str:

        language = self._get_language(
            candidate
        )

        return self.SECTION_TITLES[
            language
        ][key]

    # ---------------------------------------------------------
    # HEADER
    # ---------------------------------------------------------

    def _add_header(
        self,
        document: Document,
        candidate: CandidateProfile,
    ) -> None:

        # NAME
        if candidate.full_name:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True

            run = paragraph.add_run(
                candidate.full_name.upper()
            )

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(20)

        # CONTACT INFORMATION
        contact = document.add_paragraph()

        contact.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        contact.paragraph_format.space_before = Pt(0)
        contact.paragraph_format.space_after = Pt(9)

        first_item = True

        def add_separator():

            nonlocal first_item

            if not first_item:

                separator = contact.add_run(
                    "  |  "
                )

                separator.font.size = Pt(8.5)

            first_item = False

        # EMAIL
        if candidate.email:

            add_separator()

            run = contact.add_run(
                candidate.email
            )

            run.font.size = Pt(8.5)

        # PHONE
        if candidate.phone:

            add_separator()

            run = contact.add_run(
                candidate.phone
            )

            run.font.size = Pt(8.5)

        # LINKEDIN
        if self._is_valid_url(
            candidate.linkedin
        ):

            add_separator()

            self._add_hyperlink(
                paragraph=contact,
                text="LinkedIn",
                url=candidate.linkedin,
                font_size=9,
            )

        # GITHUB
        if self._is_valid_url(
            candidate.github
        ):

            add_separator()

            self._add_hyperlink(
                paragraph=contact,
                text="GitHub",
                url=candidate.github,
                font_size=9,
            )

        # PORTFOLIO
        if self._is_valid_url(
            candidate.portfolio
        ):

            add_separator()

            self._add_hyperlink(
                paragraph=contact,
                text="Portfolio",
                url=candidate.portfolio,
                font_size=9,
            )

    # ---------------------------------------------------------
    # URL VALIDATION
    # ---------------------------------------------------------

    def _is_valid_url(
        self,
        value: str | None,
    ) -> bool:

        if not value:
            return False

        value = value.strip().lower()

        return (
            value.startswith("http://")
            or value.startswith("https://")
        )

    # ---------------------------------------------------------
    # HYPERLINK
    # ---------------------------------------------------------

    def _add_hyperlink(
        self,
        paragraph,
        text: str,
        url: str,
        font_size: int = 9,
    ) -> None:

        part = paragraph.part

        relationship_id = part.relate_to(
            url,
            (
                "http://schemas.openxmlformats.org/"
                "officeDocument/2006/"
                "relationships/hyperlink"
            ),
            is_external=True,
        )

        hyperlink = OxmlElement(
            "w:hyperlink"
        )

        hyperlink.set(
            qn("r:id"),
            relationship_id,
        )

        new_run = OxmlElement(
            "w:r"
        )

        run_properties = OxmlElement(
            "w:rPr"
        )

        # Word hyperlink style
        run_style = OxmlElement(
            "w:rStyle"
        )

        run_style.set(
            qn("w:val"),
            "Hyperlink",
        )

        run_properties.append(
            run_style
        )

        # Font size
        size = OxmlElement(
            "w:sz"
        )

        size.set(
            qn("w:val"),
            str(font_size * 2),
        )

        run_properties.append(
            size
        )

        new_run.append(
            run_properties
        )

        text_element = OxmlElement(
            "w:t"
        )

        text_element.text = text

        new_run.append(
            text_element
        )

        hyperlink.append(
            new_run
        )

        paragraph._p.append(
            hyperlink
        )

    # ---------------------------------------------------------
    # SECTION TITLE
    # ---------------------------------------------------------

    def _add_section_title(
        self,
        document: Document,
        title: str,
    ) -> None:

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(7)
        paragraph.paragraph_format.space_after = Pt(3)

        # Prevent section title from being left alone
        # at the bottom of a page.
        paragraph.paragraph_format.keep_with_next = True

        run = paragraph.add_run(
            title.upper()
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10.5)

        # Add a thin bottom border under section title.
        paragraph_properties = (
            paragraph._p.get_or_add_pPr()
        )

        borders = OxmlElement(
            "w:pBdr"
        )

        bottom = OxmlElement(
            "w:bottom"
        )

        bottom.set(
            qn("w:val"),
            "single",
        )

        bottom.set(
            qn("w:sz"),
            "6",
        )

        bottom.set(
            qn("w:space"),
            "2",
        )

        bottom.set(
            qn("w:color"),
            "808080",
        )

        borders.append(
            bottom
        )

        paragraph_properties.append(
            borders
        )

    # ---------------------------------------------------------
    # SUMMARY
    # ---------------------------------------------------------

    def _add_summary(
        self,
        document: Document,
        candidate: CandidateProfile,
        plan: CVTailoringPlan,
    ) -> None:

        if not plan.professional_summary:
            return

        self._add_section_title(
            document,
            self._title(
                candidate,
                "summary",
            ),
        )

        paragraph = document.add_paragraph(
            plan.professional_summary
        )

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.05

    # ---------------------------------------------------------
    # SKILLS
    # ---------------------------------------------------------

    def _add_skills(
        self,
        document: Document,
        candidate: CandidateProfile,
        plan: CVTailoringPlan,
    ) -> None:

        if not plan.prioritized_skills:
            return

        self._add_section_title(
            document,
            self._title(
                candidate,
                "skills",
            ),
        )

        paragraph = document.add_paragraph(
            " • ".join(
                plan.prioritized_skills
            )
        )

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0

        for run in paragraph.runs:
            run.font.size = Pt(9)

    # ---------------------------------------------------------
    # EXPERIENCE
    # ---------------------------------------------------------

    def _add_experience(
        self,
        document: Document,
        candidate: CandidateProfile,
        plan: CVTailoringPlan,
    ) -> None:

        if not candidate.experiences:
            return

        self._add_section_title(
            document,
            self._title(
                candidate,
                "experience",
            ),
        )

        tailored_map = {
            item.experience_title: item
            for item in plan.experiences
        }

        for experience in candidate.experiences:

            heading = document.add_paragraph()

            heading.paragraph_format.space_before = Pt(3)
            heading.paragraph_format.space_after = Pt(1)
            heading.paragraph_format.keep_with_next = True

            # EXPERIENCE TITLE
            run = heading.add_run(
                experience.title
            )

            run.bold = True
            run.font.size = Pt(10)

            # ORGANIZATION
            if experience.organization:

                organization_run = heading.add_run(
                    f" — {experience.organization}"
                )

                organization_run.italic = True
                organization_run.font.size = Pt(9.5)

            tailored = tailored_map.get(
                experience.title
            )

            # Tailored bullets
            if (
                tailored
                and tailored.tailored_bullets
            ):

                for bullet in (
                    tailored.tailored_bullets
                ):

                    self._add_bullet(
                        document,
                        bullet.tailored,
                    )

            # Original bullets fallback
            else:

                for description in (
                    experience.description
                ):

                    self._add_bullet(
                        document,
                        description,
                    )

    # ---------------------------------------------------------
    # PROJECTS
    # ---------------------------------------------------------

    def _add_projects(
        self,
        document: Document,
        candidate: CandidateProfile,
        plan: CVTailoringPlan,
    ) -> None:

        if not candidate.projects:
            return

        self._add_section_title(
            document,
            self._title(
                candidate,
                "projects",
            ),
        )

        project_map = {
            project.name: project
            for project in candidate.projects
        }

        used_projects = set()

        # -----------------------------------------------------
        # Relevant projects from tailoring plan first
        # -----------------------------------------------------

        for tailored_project in plan.projects:

            original_project = project_map.get(
                tailored_project.project_name
            )

            if not original_project:
                continue

            used_projects.add(
                original_project.name
            )

            self._write_project(
                document,
                original_project,
                tailored_project,
            )

        # -----------------------------------------------------
        # Remaining original projects
        # -----------------------------------------------------

        for project in candidate.projects:

            if project.name in used_projects:
                continue

            self._write_project(
                document,
                project,
                None,
            )

    # ---------------------------------------------------------
    # PROJECT WRITER
    # ---------------------------------------------------------

    def _write_project(
        self,
        document: Document,
        project,
        tailored_project,
    ) -> None:

        # PROJECT NAME
        heading = document.add_paragraph()

        heading.paragraph_format.space_before = Pt(3)
        heading.paragraph_format.space_after = Pt(0)
        heading.paragraph_format.keep_with_next = True

        run = heading.add_run(
            project.name
        )

        run.bold = True
        run.font.size = Pt(10)

        # TECHNOLOGIES
        if project.technologies:

            technologies = document.add_paragraph()

            technologies.paragraph_format.space_before = Pt(0)
            technologies.paragraph_format.space_after = Pt(1)
            technologies.paragraph_format.keep_with_next = True

            tech_run = technologies.add_run(
                " • ".join(
                    project.technologies
                )
            )

            tech_run.italic = True
            tech_run.font.size = Pt(8.5)

        # Tailored project bullets
        if (
            tailored_project
            and tailored_project.tailored_bullets
        ):

            for bullet in (
                tailored_project.tailored_bullets
            ):

                self._add_bullet(
                    document,
                    bullet.tailored,
                )

        # Original project bullets fallback
        else:

            for description in (
                project.description
            ):

                self._add_bullet(
                    document,
                    description,
                )

    # ---------------------------------------------------------
    # EDUCATION
    # ---------------------------------------------------------

    def _add_education(
        self,
        document: Document,
        candidate: CandidateProfile,
    ) -> None:

        if not candidate.education:
            return

        self._add_section_title(
            document,
            self._title(
                candidate,
                "education",
            ),
        )

        for education in candidate.education:

            # INSTITUTION
            if education.institution:

                institution = document.add_paragraph()

                institution.paragraph_format.space_before = Pt(2)
                institution.paragraph_format.space_after = Pt(0)
                institution.paragraph_format.keep_with_next = True

                run = institution.add_run(
                    education.institution
                )

                run.bold = True
                run.font.size = Pt(10)

            # EDUCATION DETAILS
            details = []

            if education.degree:
                details.append(
                    education.degree
                )

            # Prevent duplicate:
            # Computer Engineering | Computer Engineering
            if (
                education.department
                and education.department
                != education.degree
            ):
                details.append(
                    education.department
                )

            if education.graduation_year:
                details.append(
                    str(
                        education.graduation_year
                    )
                )

            if details:

                paragraph = document.add_paragraph(
                    " | ".join(
                        details
                    )
                )

                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)

                for run in paragraph.runs:
                    run.font.size = Pt(9)

    # ---------------------------------------------------------
    # CERTIFICATIONS
    # ---------------------------------------------------------

    def _add_certifications(
        self,
        document: Document,
        candidate: CandidateProfile,
    ) -> None:

        if not candidate.certifications:
            return

        self._add_section_title(
            document,
            self._title(
                candidate,
                "certifications",
            ),
        )

        for certification in candidate.certifications:

            self._add_bullet(
                document,
                certification,
            )

    # ---------------------------------------------------------
    # LANGUAGES
    # ---------------------------------------------------------

    def _add_languages(
        self,
        document: Document,
        candidate: CandidateProfile,
    ) -> None:

        if not candidate.languages:
            return

        self._add_section_title(
            document,
            self._title(
                candidate,
                "languages",
            ),
        )

        paragraph = document.add_paragraph(
            " • ".join(
                candidate.languages
            )
        )

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)

        for run in paragraph.runs:
            run.font.size = Pt(9)

    # ---------------------------------------------------------
    # BULLET HELPER
    # ---------------------------------------------------------

    def _add_bullet(
        self,
        document: Document,
        text: str,
    ) -> None:

        if not text:
            return

        # Do not use Word's default List Bullet style.
        # Explicit formatting gives us more predictable
        # results across Word/PDF conversion and ATS parsing.
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.0

        bullet = paragraph.add_run(
            "• "
        )

        bullet.bold = True
        bullet.font.size = Pt(9)

        run = paragraph.add_run(
            text
        )

        run.font.size = Pt(9)


document_generator = DocumentGenerator()